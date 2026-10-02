from pathlib import Path
import re
import csv
import json
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
import build_proposal_docx as renderer

ROOT = Path(__file__).resolve().parents[1]
source = (ROOT / 'deliverables/CHUM_RESEARCH_PROPOSAL_KO.md').read_text(encoding='utf-8')

def section(number, new_number, title):
    match = re.search(rf'^## {number}\. .*?\n(.*?)(?=^## \d+\. |\Z)', source, re.M | re.S)
    if not match:
        raise ValueError(number)
    body = match.group(1).strip().removesuffix('---').strip()
    body = re.sub(rf'^(### ){number}\.', rf'\g<1>{new_number}.', body, flags=re.M)
    return f'## {new_number}. {title}\n\n{body}\n'

def followup_results():
    run = ROOT / 'outputs/chum_window_extension_20261001'
    status = json.loads((run / 'REVIEW_20261002.json').read_text(encoding='utf-8'))
    with (run / 'METRICS.csv').open(encoding='utf-8', newline='') as stream:
        rows = list(csv.DictReader(stream))
    keys = {(r['architecture'], r['variant'], int(r['window']), int(r['seed'])) for r in rows}
    expected = {(a, v, w, 47) for a in ('tcn', 'transformer') for v in ('F0', 'F1', 'F0-C') for w in (20, 60, 120)}
    if status['status'] != 'PASS' or status['task_count'] != 18 or len(rows) != 18 or keys != expected:
        raise ValueError('Stage 1B requires all 18 unique completed tasks')
    if not status.get('run_fingerprint') or any(r['run_fingerprint'] != status['run_fingerprint'] for r in rows):
        raise ValueError('Stage 1B metric fingerprint differs from reviewed run')
    if any(r['evaluation_role'] != 'development_benchmark' or r['test_evaluated'] != 'True' for r in rows):
        raise ValueError('Unexpected evaluation role or unevaluated task')
    with (ROOT / 'outputs/final_gate_exp1/artifacts/reinartz_split_manifest.csv').open(encoding='utf-8', newline='') as stream:
        test_runs = {r['run_index'] for r in csv.DictReader(stream) if r['split'] == 'test'}
    if len(test_runs) != 560:
        raise ValueError('Unexpected development run count')
    review = json.loads((ROOT / 'outputs/chum_harness_20261001/TRAINING_BUDGET_REVIEW.json').read_text(encoding='utf-8'))
    prior = {r['architecture']: r for r in review['rows'] if r['variant'] == 'F1'}
    table = ['| 구조 | Window | AUROC | AP (저장 필드 auprc) | 탐지 run 비율 | 탐지된 run 지연 | 미탐 패널티 포함 평균 지연 | 고장 이전 sample FPR |',
             '| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |']
    for row in sorted((r for r in rows if r['variant'] == 'F1'), key=lambda r: (r['architecture'], int(r['window']))):
        values = [row['architecture'], row['window']]
        values.extend(f"{float(row[key]):.6f}" for key in ('auroc', 'auprc', 'detected_run_ratio', 'detection_delay_detected_only', 'censored_delay_mean', 'prefault_sample_fpr'))
        table.append('| ' + ' | '.join(values) + ' |')
    return f"""### 9.5 후속 학습 길이·window 개발 실험 (2026-10-02 보충)

기존 핵심 결과와 구분하여 seed47의 후속 개발 실험을 보고한다. 1A에서는 정상 validation으로 checkpoint를 선택하며 학습 길이를 늘렸지만 F1 AUROC가 개선되지 않았다. TCN은 {prior['tcn']['auroc_10_epochs']:.6f}에서 {prior['tcn']['auroc_extended']:.6f}로, Transformer는 {prior['transformer']['auroc_10_epochs']:.6f}에서 {prior['transformer']['auroc_extended']:.6f}로 변했다. 근거는 `outputs/chum_harness_20261001/TRAINING_BUDGET_REVIEW.json`이며, 정상 예측 MSE 개선을 탐지력 개선으로 간주하지 않는다.

1B는 두 구조 × F0/F1/F0-C × window 20/60/120의 18/18 작업을 완료했다. 아래 F1 여섯 행은 `outputs/chum_window_extension_20261001/METRICS.csv`에서 직접 생성했다. 동일한 기존 test 560 runs를 개발 벤치마크로 사용하고 window 사이의 평가 시작 sample을 121로 맞췄다. 지연 단위는 sample이며, 탐지된 run의 평균 지연과 미탐에 1,401 sample의 패널티를 부여한 평균을 구분한다. 후자는 생존분석의 검열 추정량이 아니다. AP는 non-interpolated average precision이며 분류 정확도가 아니다. 평가 구간의 이상 sample 비율은 약 74.52%이다.

{chr(10).join(table)}

F1의 탐지 run 비율은 약 67.86–68.75%로 미탐 run도 약 31–32% 남았다. 고장 이전 sample FPR은 약 0.99–1.02%이다. 3연속 경보의 고장 이전 run alarm 비율이 0이라는 결과는 sample FPR이 0이라는 뜻이 아니다. 정상 validation을 checkpoint 선택과 threshold 설정에 함께 사용했으며 별도 독립 calibration 자료를 둔 실험은 아니다.

긴 window가 일관된 이득을 주지는 않았다. 이 표만으로 최적 window나 새로운 방법의 성능 우위를 확정하지 않는다. 1B는 단일 model seed47이고 이미 관찰한 개발 run을 재사용했으므로 독립적인 일반화 검증이나 새로운 통계적 유의성의 근거가 아니다. Transformer F0와 F0-C는 동일 구조이므로 독립적인 두 구조의 지지로 세지 않는다. 1B의 TCN은 1A의 두 층에서 여섯 층으로 바뀌었고 공통 평가 구간도 조정되었으므로 1A와 1B의 직접 수치 차이를 window만의 인과적 효과로 해석할 수 없다. 이 보충 결과는 학습 조건의 한계를 공개하는 자료이며 기존 채널 감사 결과의 수치를 대체하지 않는다.

"""


parts = ['''# 산업 시계열 이상 탐지에서 제어 이력의 조건부 유용성 감사

CHUM: Architecture-Robust Auditing of Control-History Utility for Industrial Time-Series Anomaly Detection

검토용 논문 초안 · 2026-09-21 작성 / 2026-10-02 후속 결과·범위 보충

기존 연구의 방법과 결과를 논문 순서로 재구성한 초안이다. 제목과 최종 주장에 대한 지도교수 승인 또는 학교 제출 완료를 뜻하지 않는다. 2026-09-21에 보존 원자료 기반 18개 검증, 224-cell 다중비교, 민감도 bootstrap을 재실행했다. 모델 재학습과 원시 telemetry 전처리는 재실행하지 않았다.

## 초록

산업 이상 탐지에서 제어 이력의 추가 효과를 평균 성능만으로 평가하면 사건별 이질성, 모델 용량, 비현실적인 입력 대치와 오경보의 영향을 구분하기 어렵다. 본 연구는 Control-History Utility Mapping(CHUM)을 통해 센서 전용 모델, 센서·제어 결합 모델, 용량을 맞춘 센서 전용 대조군을 비교하고, 정상 데이터 기반 조건부 대치와 대치 품질·오경보·반복성·모델 구조 합의를 결합한다. TEP의 28개 fault에서 TCN과 Transformer는 동일한 7개 fault의 event-level gain을 보였다. 채널 수준에서는 F4/XMV10, F19/XMV7, F19/XMV8, F25/XMV2가 두 구조의 판정 기준을 충족했다. 대치 품질을 통과한 224개 후보에 BH 보정을 적용한 결과 네 셀이 모두 q=0.05에서 유지됐다. 세 residual block 길이와 세 draw 수를 교차한 민감도 분석에서도 네 셀이 유지됐다. HAI corrected v2에서는 전역 AUROC와 AUPRC 개선 및 직접 공격된 5개 채널 셀의 score 감소가 관찰됐으나, 3개 seed와 직접 공격 조건의 제한 때문에 외부 지지로 한정한다. CHUM은 평가한 모델과 대치 분포 아래의 예측 유용성을 감사하며 물리적 원인이나 인과적 제어 효과를 식별하지 않는다.

주요어: 산업 시계열, 이상 탐지, 제어 이력, 조건부 대치, 설명 가능성, 강건성

## 1. 서론

산업 시계열에는 센서가 관측한 상태와 제어 시스템의 조작 이력이 함께 기록된다. 제어 이력은 현재 센서 상태만으로는 구분하기 어려운 동작 조건을 제공할 수 있다. 그러나 제어 변수를 입력에 추가해 평균 탐지 성능이 개선됐다는 사실만으로 어느 사건에서 어느 채널을 모델이 활용했는지는 알 수 없다. 입력 차원과 함께 변한 모델 용량도 분리해야 한다.

채널을 제거한 뒤 탐지력이 감소하는지 측정하는 접근에도 교란이 있다. 비현실적인 대치 값은 정상 구간의 score를 변화시키고 오경보를 증가시킬 수 있다. 또한 한 모델의 중요도가 다른 구조에서도 유지된다는 보장은 없다. 따라서 유용성은 효과크기 하나가 아니라 대치 품질, 정상 구간 변화, 반복 실행의 안정성과 함께 평가해야 한다.

본 연구의 질문은 어떤 사건과 제어 채널에서 이러한 통제 뒤에도 추가 탐지 정보가 유지되는가이다. 이를 위해 CHUM 평가 프로토콜을 구성하고 TEP에서 주 분석을, HAI 21.03에서 제한적 외부 검증을 수행했다. 기여는 새 검출기의 최고 성능이 아니라 용량 대조군, 정상 조건부 대치, 조건별 임계값 보정, FPR 제한과 구조 간 합의를 결합한 감사 절차 및 그 실증이다.

## 2. 관련 연구

FIT는 시간에 따른 예측 분포 변화를 통해 다변량 시계열 관측의 중요도를 평가하며 시간 의존적 분포 변화를 통제할 필요성을 다룬다. TimeSHAP은 KernelSHAP을 순차 입력으로 확장해 feature·timestep·cell 수준 attribution을 산출한다. CHUM은 산업 fault별 탐지력 손실과 오경보 제한을 함께 평가한다는 점에서 평가 대상이 다르다. [1, 2]

Hase 등의 연구는 feature removal로 생성된 분포 밖 입력이 중요도 설명과 평가를 왜곡할 수 있음을 다룬다. CHUM은 정상 데이터로 학습한 대치와 품질 검사에 이 문제의식을 반영한다. 다만 정상 구간 품질 검사를 통과했다는 사실이 이상 구간에서 정확한 조건부 분포를 복원했다는 증명은 아니다. [3]

Integrated Gradients는 sensitivity와 implementation invariance를 중심으로 제안된 attribution 방법이다. CHUM의 탐지 utility와 동일한 양을 추정하지 않으므로 채널 순위의 보조 비교로 사용한다. FIT나 TimeSHAP 대비 성능 우위를 실험한 것이 아니며 조건부 대치 자체의 최초 제안도 주장하지 않는다. [4]
''']
for args in [(5,3,'연구질문과 반증 조건'), (6,4,'CHUM 방법'), (8,5,'데이터와 분할'), (9,6,'전처리와 평가 구간'), (10,7,'모델과 용량 통제'), (11,8,'실험 설계와 평가 지표'), (12,9,'결과'), (13,10,'대치 민감도 분석')]:
    parts.append(section(*args))
parts.append('''## 11. 결과 해석과 한계

TEP의 일곱 event-level gain fault와 네 channel-level consensus cell은 서로 다른 결과 단위다. 전자는 제어 입력 전체의 추가 이득이며, 후자는 특정 채널을 조건부로 대치한 뒤 나타나는 탐지력 감소다. event gain이 있어도 imputer 품질 또는 구조 합의를 통과하지 못하면 채널 결론을 내리지 않는다. F26/XMV4의 제외는 이 구분을 보여 준다.

224-cell BH 결과는 동일 TEP test run 분포에 대한 분석이다. run별 두 구조와 다섯 seed의 평균 효과를 사용한 단측 sign-flip 검정은 귀무가설 아래 run-cluster 부호 대칭을 가정한다. 공유 모델과 데이터가 만드는 의존성도 남는다. BH 보정을 모든 의존 구조에서의 오류율 보장으로 해석하거나, 5개 모델 seed를 5개 독립 데이터셋으로 해석하지 않는다. seed·architecture·민감도 설정의 표결을 독립으로 곱한 결합 확률은 정식 유의확률로 사용하지 않는다.

민감도 결과는 선택된 네 셀에 대한 post-hoc 강건성 분석이다. Git 이력에서 규칙의 사전등록을 입증할 수 없으며, 9개 설정의 반복 성공은 독립적인 9개 연구가 아니다. 다른 대치 모델·운전 조건까지 강건성을 보장하지 않는다.

정상 데이터 기반 대치는 관찰적 분석이다. 다른 채널도 fault의 영향을 받는 경우 대치값에 이상 정보가 잔존할 수 있고, 조건부 모델 오차가 효과에 영향을 줄 수 있다. 품질 gate와 FPR 제한은 이를 점검하지만 교란의 제거를 증명하지 않는다. 선택된 채널을 물리적 고장 원인 또는 최적 제어 대상으로 해석하지 않는다.

HAI에서는 AUROC·AUPRC의 3-seed t 구간은 양수지만 eTaF1 구간은 0을 포함한다. 직접 공격된 제어 채널의 utility는 공격받지 않은 제어 이력의 일반적 가치와 구별해야 한다. TEP와 HAI의 수치도 서로 다른 사건 정의와 평가 조건에서 얻었으므로 효과크기를 단순 비교하지 않는다.

## 12. 재현성과 연구 무결성

2026-09-21 재검증은 원 증거 커밋 bc6166f792e3faceb060a50de6219dddca393fd6에서 필요한 표·manifest 25개를 임시 폴더에 읽어 기존 validator를 실행했다. 파일별 SHA-256과 validator 해시는 outputs/closeout_validation/ARCHIVE_PROVENANCE.json에 있다. 총 18개 검증이 통과했다. 무효화된 HAI v1에서는 무효화 상태만 확인하며 성능 수치를 사용하지 않았다.

별도로 200,000회 sign-flip 기반 BH 표를 재계산해 보존된 224행 전체와 비교했고 수치가 일치했다. 민감도 분석은 현재 보존된 fault/run 표에서 2,000회 bootstrap을 재실행했으며 4/4 consensus PASS가 일치했다. HAI 전역 지표의 t 구간도 원 seed별 지표로 재계산했다.

이 절의 재현은 결과 표에서 분석을 다시 수행하는 재현이다. 모델 재학습, 원시 telemetry 전처리, 기존 G3 신뢰구간의 새 추정은 포함하지 않는다. HAI overlap 제거와 역할 분류는 보존 manifest를 점검한 것이다. 재실행 명령과 상세 범위는 CHUM_CLOSEOUT_KO.md를 따른다.

## 13. 결론

평가한 산업 시계열에서 제어 이력의 이득은 사건과 채널에 따라 달랐다. CHUM은 이 차이를 모델 용량, 정상 조건부 대치의 품질, 오경보 변화, 반복 실행 및 모델 구조 합의를 함께 고려해 측정했다. TEP에서는 네 fault–channel 셀이 두 구조와 민감도 설정에서 유지됐고, HAI에서는 제한적인 외부 지지를 얻었다. 이는 제어 이력의 예측 유용성에 대한 근거이며 인과 또는 root-cause 식별의 근거는 아니다. 향후 확장은 공격받지 않은 제어 이력과 새로운 운전 조건에서 같은 평가 절차가 유지되는지 검증하는 것이다.

## 참고문헌

[1] Tonekaboni, S., Joshi, S., Campbell, K., Duvenaud, D. K., and Goldenberg, A. (2020). What went wrong and when? Instance-wise feature importance for time-series black-box models. NeurIPS 33. https://proceedings.neurips.cc/paper/2020/hash/08fa43588c2571ade19bc0fa5936e028-Abstract.html

[2] Bento, J., Saleiro, P., Cruz, A. F., Figueiredo, M. A. T., and Bizarro, P. (2021). TimeSHAP: Explaining Recurrent Models through Sequence Perturbations. KDD. https://arxiv.org/abs/2012.00073

[3] Hase, P., Xie, H., and Bansal, M. (2021). The Out-of-Distribution Problem in Explainability and Search Methods for Feature Importance Explanations. NeurIPS 34. https://proceedings.neurips.cc/paper/2021/hash/1def1713ebf17722cbe300cfc1c88558-Abstract.html

[4] Sundararajan, M., Taly, A., and Yan, Q. (2017). Axiomatic Attribution for Deep Networks. ICML, PMLR 70:3319–3328. https://proceedings.mlr.press/v70/sundararajan17a.html

[5] Reinartz, C., Kulahci, M., and Ravn, O. (2021). An extended Tennessee Eastman simulation dataset for fault-detection and decision support systems. Computers & Chemical Engineering, 149, 107281. https://doi.org/10.1016/j.compchemeng.2021.107281

[6] AIRI Institute. FDDBenchmark, reinartz_tep distribution. https://github.com/AIRI-Institute/fddbenchmark/blob/main/README.md (접근: 2026-09-21).

[7] ICS Dataset. HIL-based Augmented ICS (HAI) Security Dataset, release 21.03. https://github.com/icsdataset/hai (접근: 2026-09-21).

[8] eTaPR maintainers. eTaPR evaluation implementation. https://github.com/wshw4ng/eTaPR (접근: 2026-09-21; 기존 프로젝트의 saurf4ng/eTaPR 주소가 이 주소로 연결됨).

문헌 [1–4]의 서지 및 요약은 원 출판처 또는 저자 arXiv 초록으로 확인했다. TEP의 upstream 연구 [5]와 실제 사용한 FDDBenchmark 배포 계층 [6]을 구분한다. 본 연구의 2,800 runs·52개 변수는 로컬 배포본 감사 결과이며 upstream 전체 데이터의 모든 모드·변수를 사용했다는 뜻이 아니다. 배포본의 세부 변환 이력이 완전히 입증된 것은 아니다. HAI와 평가 구현의 출처는 [7, 8]이다. 기존 실행에 사용된 upstream commit까지 고정한 원시 데이터 재현은 이번 재검증 범위 밖이다.
''')
text = '\n\n'.join(parts)
# Avoid carrying over a discussion of instantaneous gradients: IG integrates along a path.
text = text.replace('즉 순간 score gradient는', '즉 경로 적분 기반 attribution은')
text = text.replace('local sensitivity와 conditional necessity', '경로 기반 score attribution과 조건부 탐지 utility')
text = text.replace('공격을 직접 막는다', '해당 설정 범위에서의 강건성을 지지한다')
text = text.replace('반론을 직접 차단한다', '가능성이 평가한 설정 범위에서는 관찰되지 않았음을 보여 준다')
text = text.replace('모든 fault에서 두 architecture 모두 5/5 seed가 같은 방향이었다.', '해당 일곱 GAIN fault에서 두 architecture 모두 5/5 seed가 같은 방향이었다.')
text = text.replace('5. test label은 metric 계산에만 사용하고 training·scaling·imputation·threshold selection에 사용하지 않는다.', '5. test label은 metric 계산에만 사용하고 training·scaling·imputation·threshold selection에 사용하지 않는다. 이 구분이 미사용 평가를 뜻하지는 않는다. 기존 TEP test는 후속 방법 개발에서 반복 관찰되었으므로 후속 결과에서는 개발 벤치마크이며 새로운 holdout이 아니다.')
text = text.replace('## 10. 대치 민감도 분석', followup_results() + '## 10. 대치 민감도 분석')
md = ROOT / 'deliverables/CHUM_THESIS_DRAFT_KO.md'
md.write_text(text, encoding='utf-8')
doc = Document()
renderer.style_document(doc)
renderer.configure_page(doc.sections[0])
lines = text.splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip() or line == '---':
        i += 1
        continue
    if line.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[i+1]):
        table = [line, lines[i+1]]
        i += 2
        while i < len(lines) and lines[i].startswith('|'):
            table.append(lines[i])
            i += 1
        renderer.add_markdown_table(doc, table)
        if 'Window' in table[0]:
            # Keep the six-row follow-up comparison on one review page.
            for row in doc.tables[-1].rows[:-1]:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
        continue
    heading = re.match(r'^(#{1,4})\s+(.*)$', line)
    if heading:
        hashes, title = heading.groups()
        doc.add_heading(title, level=len(hashes)-1)
    else:
        bullet = re.match(r'^-\s+(.*)$', line)
        numbered = re.match(r'^\d+\.\s+(.*)$', line)
        if bullet:
            paragraph = doc.add_paragraph(style='List Bullet')
            line = bullet.group(1)
        elif numbered:
            # Preserve explicit Markdown numbering; Word otherwise continues
            # a previous section's List Number sequence.
            paragraph = doc.add_paragraph()
        else:
            paragraph = doc.add_paragraph()
        renderer.add_inline(paragraph, line.removeprefix('> '))
    i += 1
footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = 1  # centered review page number
page_field = OxmlElement('w:fldSimple')
page_field.set(renderer.qn('w:instr'), 'PAGE')
footer._p.append(page_field)

doc.core_properties.title = 'CHUM 검토용 논문 초안'
doc.core_properties.subject = '검증된 실험 결과와 주장 범위'
doc.save(ROOT / 'deliverables/CHUM_THESIS_DRAFT_KO.docx')
print('Wrote Markdown and Word draft:', len(text), 'characters,', len(doc.tables), 'tables')

