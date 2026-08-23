from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "deliverables" / "CHUM_RESEARCH_PROPOSAL_KO.md"
OUTPUT = ROOT / "deliverables" / "CHUM_RESEARCH_PROPOSAL_KO.docx"
ASSETS = ROOT / "deliverables" / "ppt_assets"

NAVY = "17324D"
BLUE = "2F6BFF"
TEAL = "278A78"
CYAN = "3AA7B8"
AMBER = "E4A11B"
LIGHT = "EEF3F8"
PALE_BLUE = "F4F7FF"
PALE_AMBER = "FFF8E8"
MID = "C8D5E3"
GRAY = "5D6B78"
DARK = "12202F"
WHITE = "FFFFFF"
FONT = "Malgun Gothic"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size=None, color=None, bold=None, italic=None, name=FONT) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_border(paragraph, color=BLUE, size="12", space="8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_field(paragraph) -> None:
    paragraph.add_run("Page ")
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, end])


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    configure_page(section)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.0)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.28

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 9),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, NAVY, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.12

    for list_name in ["List Bullet", "List Number"]:
        style = doc.styles[list_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(9.8)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(3.5)
        style.paragraph_format.line_spacing = 1.18

    if "Figure Caption" not in [style.name for style in doc.styles]:
        caption = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = doc.styles["Figure Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False


def configure_page(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(2.15)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.75)
    section.different_first_page_header_footer = False


def add_body_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header
    p = header.paragraphs[0]
    p.text = ""

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""


def add_inline(paragraph, text: str, default_size=10.2, default_color=DARK) -> None:
    pattern = re.compile(r"(\*\*.*?\*\*|\*[^*\n]+?\*|`.*?`|https?://\S+)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=default_size, color=default_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, color=NAVY, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=default_size, color=default_color, italic=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=default_size - 0.3, color=BLUE, bold=True)
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
        else:
            run = paragraph.add_run(token.rstrip(".,);"))
            set_run_font(run, size=8.5, color=BLUE)
            run.underline = True
            trailing = token[len(token.rstrip(".,);")):]
            if trailing:
                tr = paragraph.add_run(trailing)
                set_run_font(tr, size=default_size, color=default_color)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=default_size, color=default_color)


def add_cover(doc: Document) -> None:
    cover_section = doc.sections[0]
    cover_section.header.is_linked_to_previous = False
    cover_section.footer.is_linked_to_previous = False
    hp = cover_section.header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run("SUNGKYUNKWAN UNIVERSITY  ·  SUPERINTELLIGENCE LABORATORY")
    set_run_font(hr, size=8.5, color=GRAY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MASTER'S THESIS RESEARCH PROPOSAL")
    set_run_font(r, size=10.5, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("CHUM")
    set_run_font(r, size=28, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("신뢰 가능한 산업 시계열 이상 탐지를 위한\n아키텍처 강건형 제어 이력 유용성 감사")
    set_run_font(r, size=20, color=DARK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("Architecture-Robust Auditing of Control-History Utility\nfor Reliable Industrial Time-Series Anomaly Detection")
    set_run_font(r, size=11.5, color=GRAY, italic=True)

    meta = doc.add_table(rows=2, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    widths = [Inches(1.0), Inches(2.25), Inches(1.0), Inches(2.25)]
    values = [
        ("검토 대상", "추현승 교수", "현재 상태", "MANUSCRIPT-READY"),
        ("작성일", "2026-08-23", "필수 실험", "완료 · backlog 없음"),
    ]
    for row_index, row_values in enumerate(values):
        for col_index, value in enumerate(row_values):
            cell = meta.cell(row_index, col_index)
            cell.width = widths[col_index]
            set_cell_margins(cell, top=90, bottom=90)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            if col_index % 2 == 0:
                set_run_font(r, size=8.5, color=GRAY, bold=True)
                set_cell_shading(cell, LIGHT)
            else:
                set_run_font(r, size=9.5, color=NAVY, bold=True)
                set_cell_shading(cell, WHITE)

    doc.add_paragraph().paragraph_format.space_after = Pt(7)
    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = callout.cell(0, 0)
    set_cell_shading(cell, PALE_AMBER)
    set_cell_margins(cell, top=180, bottom=180, start=180, end=180)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    add_inline(
        p,
        "결론: 4개 locked event–channel이 TCN·Transformer와 9개 replacement 설정에서 모두 재현됐다. 현재 claim 범위에서 남은 필수 실험은 없다.",
        default_size=11.5,
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    metrics = doc.add_table(rows=2, cols=4)
    metrics.alignment = WD_TABLE_ALIGNMENT.CENTER
    metrics.autofit = False
    labels = ["LOCKED CELLS", "DATASETS", "ARCHITECTURES", "EVIDENCE QA"]
    values = ["4 / 4", "2", "3", "18 / 18"]
    for col in range(4):
        for row in range(2):
            cell = metrics.cell(row, col)
            cell.width = Inches(1.625)
            set_cell_margins(cell, top=65, bottom=65)
            set_cell_shading(cell, PALE_BLUE if col != 3 else PALE_AMBER)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(labels[col] if row == 0 else values[col])
            set_run_font(r, size=7.5 if row == 0 else 16, color=GRAY if row == 0 else (BLUE if col != 3 else AMBER), bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Decision requested  ·  연구질문 확정  |  CHUM을 주 기여로 확정  |  본문 집필 시작")
    set_run_font(r, size=9.5, color=NAVY, bold=True)
    set_paragraph_border(p, color=BLUE, size="10", space="9")


def add_document_flow(doc: Document) -> None:
    h = doc.add_heading("문서 흐름", level=1)
    h.paragraph_format.space_before = Pt(0)
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    flows = [
        ("01  WHY", "교수·연구실 최근 경향 → 문제의식 → 연구 공백"),
        ("02  WHAT", "연구질문·가설 → CHUM의 조작적 정의와 기여"),
        ("03  HOW", "Dataset → 전처리 → 모델 → 실험·판정 gate"),
        ("04  PROOF", "TEP → IG → HAI corrected v2 → 최종 민감도"),
        ("05  DECIDE", "의의·한계 → 발전 가능성 → 집필 구조와 결정 요청"),
    ]
    for i, (label, detail) in enumerate(flows):
        left, right = table.rows[i].cells
        left.width, right.width = Inches(1.4), Inches(5.1)
        set_cell_shading(left, NAVY if i < 4 else AMBER)
        set_cell_shading(right, LIGHT)
        for cell in (left, right):
            set_cell_margins(cell, top=120, bottom=120, start=140, end=140)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        lp, rp = left.paragraphs[0], right.paragraphs[0]
        lp.paragraph_format.space_after = rp.paragraph_format.space_after = Pt(0)
        lr, rr = lp.add_run(label), rp.add_run(detail)
        set_run_font(lr, size=9, color=WHITE, bold=True)
        set_run_font(rr, size=10.5, color=DARK, bold=(i == 4))
    doc.add_paragraph()
    p = doc.add_paragraph()
    add_inline(p, "이 문서는 결론을 먼저 제시하고, 각 주장에 대응하는 통제와 raw-evidence 판정을 순서대로 연결한다.")
    p.paragraph_format.space_after = Pt(0)


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_AMBER)
    set_cell_margins(cell, top=130, bottom=130, start=150, end=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, text, default_size=10.7)


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        parts = [part.strip() for part in line.strip().strip("|").split("|")]
        rows.append(parts)
    if len(rows) < 2:
        return
    data_rows = [rows[0]] + rows[2:]
    cols = max(len(row) for row in data_rows)
    table = doc.add_table(rows=len(data_rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    available = 6.48
    for row_index, values in enumerate(data_rows):
        prevent_row_split(table.rows[row_index])
        for col_index in range(cols):
            cell = table.cell(row_index, col_index)
            cell.width = Inches(available / cols)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 or col_index > 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(
                p,
                values[col_index] if col_index < len(values) else "",
                default_size=7.6 if cols >= 6 else (8.1 if cols >= 5 else 8.7),
            )
            for run in p.runs:
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(NAVY)
    table.rows[0].height = None
    if len(data_rows) <= 5:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def new_numbering_id(doc: Document) -> int:
    """Create a fresh single-level decimal list so every Markdown block restarts at 1."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_format = OxmlElement("w:numFmt")
    num_format.set(qn("w:val"), "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "540")
    indentation.set(qn("w:hanging"), "280")
    paragraph_properties.extend([tabs, indentation])
    level.extend([start, num_format, level_text, level_justification, paragraph_properties])
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_element])
    p_pr.append(num_pr)


FIGURES = {
    "2. 추현승 교수·연구실 최근 연구 경향과 본 연구의 적합성": (
        "06_PROFESSOR_FIT_MAP.png",
        "그림 1. 공식 연구실 2024–2026 publication 제목 코딩과 CHUM의 제안 위치",
    ),
    "6. 해결 방식: CHUM": (
        "01_CHUM_METHOD_PIPELINE.png",
        "그림 2. CHUM의 다섯 단계와 해석 경계",
    ),
    "8. Dataset과 split": (
        "02_DATASET_AND_SPLIT_PROTOCOL.png",
        "그림 3. TEP·HAI 데이터 환경과 누수 방지 protocol",
    ),
    "12.2 TEP channel-level consensus": (
        "03_TEP_ARCHITECTURE_CONSENSUS.png",
        "그림 4. TEP locked cell의 두 architecture ΔAUROC",
    ),
    "12.4 HAI 21.03 corrected v2": (
        "05_HAI_EXTERNAL_SUPPORT.png",
        "그림 5. HAI corrected v2의 전역 및 targeted-channel 외부 지지",
    ),
    "13.2 결과": (
        "04_PRIMARY_SENSITIVITY_RANGES.png",
        "그림 6. 3×3 replacement sensitivity의 architecture–cell별 ΔAUROC 범위",
    ),
    "14. 연구 무결성과 재현성": (
        "07_EVIDENCE_SCORECARD.png",
        "그림 7. 논문 집필 전 핵심 반론과 대응 evidence scorecard",
    ),
}


PAGE_BREAK_SECTIONS = {
    "2. 추현승 교수·연구실 최근 연구 경향과 본 연구의 적합성",
    "13. 새로 완료한 최종 민감도 실험",
}


def render_markdown(doc: Document) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 0."))
    lines = lines[start:]
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line or line == "---":
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[index + 1]):
            table_lines = [line, lines[index + 1]]
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(doc, table_lines)
            continue
        if line.startswith("> "):
            add_callout(doc, line[2:])
            index += 1
            continue
        heading_match = re.match(r"^(#{2,4})\s+(.*)$", line)
        if heading_match:
            hashes, title = heading_match.groups()
            level = min(len(hashes) - 1, 3)
            h = doc.add_heading(title, level=level)
            if title in PAGE_BREAK_SECTIONS and len(doc.paragraphs) > 2:
                h.paragraph_format.page_break_before = True
            if level == 1:
                set_paragraph_border(h, color=MID, size="6", space="4")
            if title in FIGURES:
                name, caption = FIGURES[title]
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(ASSETS / name), width=Inches(6.1))
                cp = doc.add_paragraph(caption, style="Figure Caption")
                cp.paragraph_format.keep_with_next = False
            index += 1
            continue
        bullet = re.match(r"^-\s+(.*)$", line)
        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1), default_size=9.9)
            index += 1
            continue
        if numbered:
            num_id = new_numbering_id(doc)
            while index < len(lines):
                current = re.match(r"^\d+\.\s+(.*)$", lines[index].rstrip())
                if not current:
                    break
                p = doc.add_paragraph(style="List Number")
                apply_numbering(p, num_id)
                add_inline(p, current.group(1), default_size=9.9)
                index += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, line)
        index += 1


def main() -> None:
    doc = Document()
    style_document(doc)
    add_cover(doc)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_page(body_section)
    add_body_header_footer(body_section)
    add_document_flow(doc)
    doc.add_page_break()
    render_markdown(doc)

    properties = doc.core_properties
    properties.title = "CHUM 연구 기획서"
    properties.subject = "추현승 교수 검토용 석사논문 연구 기획"
    properties.author = "Thesis-Orchestrator"
    properties.keywords = "CHUM, industrial time series, anomaly detection, control history, trustworthy AI"
    properties.comments = "Generated from CHUM_RESEARCH_PROPOSAL_KO.md; narrative_proposal preset with proposal_centerpiece cover."
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
